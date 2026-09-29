"""选产品分页合成响应回归；这些响应不是官网抓包样本。"""

import asyncio
import json
import time
from dataclasses import replace
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from opscli.seller_sprite.domain.exceptions import (
    SellerSpriteApiError,
    SellerSpriteConfigError,
    SellerSpriteTaskTimeoutError,
)
from opscli.seller_sprite.domain.models import SellerSpriteScenarioRequest
from opscli.seller_sprite.services.product_pagination import (
    ProductPagination,
    parse_product_page,
)


def response(total, page, size=100):
    """构造可重复、带唯一子ASIN的合成分页响应。"""
    return {
        "code": "OK",
        "data": {
            "total": total,
            "items": [
                {
                    "asin": f"B{i:09d}",
                    "title": f"测试商品{i}",
                    "price": 12.34,
                    "monthlySales": i,
                }
                for i in range((page - 1) * size, min(page * size, total))
            ],
        },
    }


def collector(tmp_path, **kwargs):
    return ProductPagination(
        root=tmp_path / "attempt-1",
        task_root=tmp_path,
        account_key="account-1",
        request_identity={"keywords": "LOFT"},
        payload={
            "keywords": "LOFT",
            "page": 1,
            "size": 100,
            "order": {"field": "total_units", "desc": True},
        },
        timeout_seconds=600,
        **kwargs,
    )


@pytest.mark.parametrize(
    "total,pages,rows,coverage",
    [
        (0, 1, 0, "matched_range"),
        (21, 1, 21, "matched_range"),
        (100, 1, 100, "matched_range"),
        (121, 2, 121, "matched_range"),
        (2000, 20, 2000, "matched_range"),
        (4830, 20, 2000, "source_limited"),
    ],
)
def test_accessible_range(tmp_path, total, pages, rows, coverage):
    instance = collector(tmp_path)
    calls = []

    async def fetch(payload, root):
        calls.append(payload)
        return response(total, payload["page"])

    asyncio.run(instance.collect(fetch))
    assert [item["page"] for item in calls] == list(range(1, pages + 1))
    assert all({**item, "page": 1} == instance.state["payload"] for item in calls)
    assert instance.summary()["coverage"] == coverage
    assert instance.summary()["collection_complete"] is True
    assert len(list(instance.iter_rows())) == rows
    manifest = json.loads((instance.root / "raw.json").read_text(encoding="utf-8"))
    assert "items" not in json.dumps(manifest)


@pytest.mark.parametrize("change", ["duplicate", "total", "short"])
def test_inconsistent_pages_not_complete(tmp_path, change):
    instance = collector(tmp_path)

    async def fetch(payload, root):
        result = response(250, payload["page"])
        if payload["page"] == 2:
            if change == "duplicate":
                result["data"]["items"][0]["asin"] = "B000000000"
            elif change == "total":
                result["data"]["total"] = 249
            else:
                result["data"]["items"] = []
        return result

    asyncio.run(instance.collect(fetch))
    assert instance.summary()["coverage"] == "unverified"
    assert instance.summary()["collection_complete"] is False
    assert instance.summary()["pages_completed"] == 2


def test_resume_skips_committed_pages_and_validates_files(tmp_path):
    saved = []
    first = collector(
        tmp_path, save_checkpoint=lambda cp: saved.append(json.loads(json.dumps(cp)))
    )

    async def interrupted(payload, root):
        if payload["page"] == 2:
            raise RuntimeError("模拟进程退出")
        return response(121, 1)

    with pytest.raises(RuntimeError):
        asyncio.run(first.collect(interrupted))
    second = collector(tmp_path, checkpoint=saved[-1])
    second.root = tmp_path / "attempt-2"
    calls = []

    async def resume(payload, root):
        calls.append(payload["page"])
        return response(121, payload["page"])

    asyncio.run(second.collect(resume))
    assert calls == [2]
    assert len(list(second.iter_rows())) == 121
    assert second.state["deadline"] == first.state["deadline"]
    (tmp_path / first.state["pages"][0]["path"]).write_text("{}", encoding="utf-8")
    with pytest.raises(SellerSpriteConfigError, match="校验失败"):
        collector(tmp_path, checkpoint=second.state)


def test_resume_rejects_account_query_and_expired_budget(tmp_path):
    instance = collector(tmp_path)

    async def fetch(payload, root):
        return response(21, 1)

    asyncio.run(instance.collect(fetch))
    checkpoint = json.loads(json.dumps(instance.state))
    checkpoint["account_key"] = "other"
    with pytest.raises(SellerSpriteConfigError, match="账号"):
        collector(tmp_path, checkpoint=checkpoint)
    checkpoint["identity"] = "changed"
    with pytest.raises(SellerSpriteConfigError, match="条件"):
        collector(tmp_path, checkpoint=checkpoint)
    instance.state["deadline"] = time.time() - 1
    with pytest.raises(SellerSpriteTaskTimeoutError):
        instance.remaining()


def test_retry_budget_survives_recovery(tmp_path, monkeypatch):
    async def no_sleep(seconds):
        return None

    monkeypatch.setattr(
        "opscli.seller_sprite.services.product_pagination.asyncio.sleep", no_sleep
    )
    instance = collector(tmp_path)
    calls = []

    async def fail(payload, root):
        calls.append(payload["page"])
        raise SellerSpriteApiError("临时故障", status_code=503)

    with pytest.raises(SellerSpriteApiError):
        asyncio.run(instance.collect(fail))
    assert calls == [1, 1, 1]
    resumed = collector(tmp_path, checkpoint=instance.state)
    with pytest.raises(SellerSpriteConfigError, match="预算"):
        asyncio.run(resumed.collect(fail))
    assert calls == [1, 1, 1]


@pytest.mark.parametrize(
    "data",
    [
        {},
        {"items": [], "total": True},
        {"items": [], "total": -1},
        {"pager": {"items": [], "total": 0}},
        {"items": [None], "total": 1},
    ],
)
def test_does_not_guess_other_scenario_response(data):
    with pytest.raises(SellerSpriteConfigError):
        parse_product_page({"data": data})


@pytest.mark.parametrize(
    "overrides",
    [
        {"scenario": "listing-analysis"},
        {"export_format": "json"},
        {"page_size": 20},
        {"params": {"size": 60}},
        {"params": {"pageSize": 100, "size": 20}},
        {"params": {"startPage": 2}},
        {"params": {"page": True}},
        {"params": {"page": 1, "pageNum": 2}},
        {"mode": "api-direct"},
        {"collection_scope": "everything"},
    ],
)
def test_all_request_validation(overrides):
    with pytest.raises(SellerSpriteConfigError):
        SellerSpriteScenarioRequest(
            **{
                "scenario": "product-research",
                "site": "US",
                "period": "30d",
                "collection_scope": "all",
                **overrides,
            }
        )


@pytest.mark.parametrize(
    "scenario",
    [
        "product-research",
        "competitor-lookup",
        "keyword-miner",
        "keyword-reverse",
        "traffic-source",
        "association-traffic",
        "aba-research",
        "keyword-research",
        "market-research",
    ],
)
def test_all_request_validation_allows_verified_list_scenarios(scenario):
    request = SellerSpriteScenarioRequest(
        scenario=scenario,
        site="US",
        period="30d",
        collection_scope="all",
        mode="browser-route",
        export_format="xlsx",
        page_size=100,
    )
    assert request.collection_scope == "all"


@pytest.mark.parametrize(
    "scenario",
    [
        "aba-reverse",
        "branddb",
        "listing-analysis",
        "keyword-comparison",
        "traffic-extend",
        "keyword-conversion-rate",
        "real-time-bidding",
    ],
)
def test_all_request_validation_rejects_non_list_scenarios(scenario):
    with pytest.raises(SellerSpriteConfigError, match="已核验的列表场景"):
        SellerSpriteScenarioRequest(
            scenario=scenario,
            site="US",
            period="30d",
            collection_scope="all",
        )


def test_streaming_export_matches_single_page(tmp_path):
    from opscli.seller_sprite.export.xlsx import (
        export_product_pages_to_xlsx,
        export_rows_to_xlsx,
    )

    rows = response(21, 1)["data"]["items"]
    export_rows_to_xlsx(
        rows=rows, output_path=tmp_path / "old.xlsx", scenario="product-research"
    )
    export_product_pages_to_xlsx(
        rows=iter(rows),
        output_path=tmp_path / "new.xlsx",
        site="US",
        period="30d",
        params={},
    )
    before = load_workbook(tmp_path / "old.xlsx")
    after = load_workbook(tmp_path / "new.xlsx")
    try:
        assert before.sheetnames == after.sheetnames
        assert list(before.active.values) == list(after.active.values)
        assert [[c.number_format for c in row] for row in before.active] == [
            [c.number_format for c in row] for row in after.active
        ]
        assert after.active.freeze_panes == "A2"
    finally:
        before.close()
        after.close()


def test_manager_all_result_is_bounded_and_can_be_ingested(tmp_path, monkeypatch):
    from opscli.seller_sprite.collection_storage_parser import (
        SellerSpriteCollectionParser,
    )
    from opscli.seller_sprite.config import SellerSpriteSettings
    from opscli.seller_sprite.services import api_manager as module
    from opscli.shared.collection_storage.models import CollectionSubmission
    from tests.seller_sprite.test_api_manager import (
        DummyAccountProvider,
        DummyApiClient,
    )

    monkeypatch.setattr(module, "SellerSpriteApiClient", DummyApiClient)
    monkeypatch.setattr(module, "_upload_export_if_enabled", lambda **kw: None)
    calls = []

    async def browser(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(
            response=response(121, kwargs["payload"]["page"]), warnings=[]
        )

    monkeypatch.setattr(module, "_run_browser_route_request", browser)
    manager = module.SellerSpriteApiManager(
        settings=SellerSpriteSettings(output_dir=tmp_path),
        account_provider=DummyAccountProvider(),
    )
    result = asyncio.run(
        manager.run(
            SellerSpriteScenarioRequest(
                scenario="product-research",
                site="US",
                period="30d",
                collection_scope="all",
                mode="browser-route",
                job_id="all-test",
                params={"keywords": "LOFT"},
            )
        )
    )
    assert result.row_count == 121
    assert result.data == []
    assert len(result.preview) == 10
    assert result.collection["collection_complete"] is True
    assert [call["payload"]["page"] for call in calls] == [1, 2]
    parsed = SellerSpriteCollectionParser().parse(
        CollectionSubmission(
            source_system="seller_sprite",
            source_job_id=result.job_id,
            producer_service="test",
            scenario="product-research",
            site="US",
            data_environment="debug",
            ingestion_mode="live",
            result_path=result.result_path,
        )
    )
    assert len(list(parsed.datasets[0].records)) == 121


def test_queue_checkpoint_requires_live_owner_generation_account(tmp_path):
    from opscli.seller_sprite.services.task_queue_store import (
        SellerSpriteTaskQueueStore,
    )

    store = SellerSpriteTaskQueueStore(db_path=tmp_path / "queue.db")
    request = SellerSpriteScenarioRequest(
        scenario="product-research",
        site="US",
        period="30d",
        job_id="job",
        collection_scope="all",
    )
    store.enqueue(request=request, queue_scope="test", root_dir=tmp_path)
    claim = store.claim_next_generic_for_account(
        queue_scope="test",
        account_key="account-1",
        assigned_account="a",
        worker_key="worker",
        execution_owner="owner",
    )
    args = {
        "job_id": "job",
        "account_key": "account-1",
        "assignment_generation": claim["assignment_generation"],
        "execution_owner": "owner",
        "checkpoint": {"collection": {"coverage": "partial"}},
    }
    assert store.save_product_checkpoint(**args)
    for change in (
        {"account_key": "other"},
        {"assignment_generation": 0},
        {"execution_owner": "other"},
    ):
        assert not store.save_product_checkpoint(**{**args, **change})
    assert store.get_status("job")["collection"]["coverage"] == "partial"
    with store._connect() as connection:
        connection.execute(
            "UPDATE seller_sprite_task_queue SET lease_expires_at = '2000-01-01'"
        )
    assert not store.save_product_checkpoint(**args)


def test_all_cache_identity_and_submission_are_isolated(tmp_path):
    from opscli.seller_sprite.collection_storage_integration import (
        SellerSpriteCollectionSubmitter,
        build_seller_sprite_cache_identity,
    )
    from opscli.seller_sprite.domain.models import SellerSpriteScenarioResult

    page = SellerSpriteScenarioRequest(
        scenario="product-research", site="US", period="30d"
    )
    full = replace(page, collection_scope="all")
    assert build_seller_sprite_cache_identity(
        page, account_route=None, requested_account_key=None
    ) != build_seller_sprite_cache_identity(
        full, account_route=None, requested_account_key=None
    )
    submissions = []
    runtime = SimpleNamespace(
        settings=SimpleNamespace(data_environment="debug"),
        submit=lambda item: submissions.append(item) or True,
    )
    result = SellerSpriteScenarioResult.empty(
        job_id="job",
        scenario="product-research",
        site="US",
        period="30d",
        root_dir=tmp_path,
        params_path=tmp_path / "params.json",
        raw_path=tmp_path / "raw.json",
        result_path=tmp_path / "result.json",
    )
    assert SellerSpriteCollectionSubmitter(runtime)(
        request=full, result=result, status={"state": "succeeded"}
    )
    assert submissions[0].cache_key is None
    assert submissions[0].cache_scope is None


def test_recovery_claim_keeps_page_account_and_rejects_old_writer(tmp_path):
    from opscli.seller_sprite.services.task_queue_store import (
        SellerSpriteTaskQueueStore,
    )

    store = SellerSpriteTaskQueueStore(db_path=tmp_path / "queue.db")
    request = SellerSpriteScenarioRequest(
        scenario="product-research",
        site="US",
        period="30d",
        job_id="job",
        collection_scope="all",
    )
    store.enqueue(request=request, queue_scope="test", root_dir=tmp_path)
    first = store.claim_next_generic_for_account(
        queue_scope="test",
        account_key="a",
        assigned_account="a",
        worker_key="w",
        execution_owner="old",
    )
    args = {
        "job_id": "job",
        "account_key": "a",
        "assignment_generation": first["assignment_generation"],
        "execution_owner": "old",
        "checkpoint": {
            "pages": [{"page": 1}],
            "account_key": "a",
            "deadline": time.time() + 600,
        },
    }
    assert store.save_product_checkpoint(**args)
    assert store.reset_running_tasks() == 1
    assert (
        store.claim_next_generic_for_account(
            queue_scope="test",
            account_key="b",
            assigned_account="b",
            worker_key="w2",
            execution_owner="new",
        )
        is None
    )
    second = store.claim_next_generic_for_account(
        queue_scope="test",
        account_key="a",
        assigned_account="a",
        worker_key="w2",
        execution_owner="new",
    )
    assert second["assignment_generation"] > first["assignment_generation"]
    assert not store.save_product_checkpoint(**args)


@pytest.mark.parametrize(
    "dedicated,fail_page_two", [(False, False), (True, False), (False, True)]
)
def test_scheduler_collects_pages_or_preserves_partial_without_failover(
    tmp_path, monkeypatch, dedicated, fail_page_two
):
    from opscli.seller_sprite.config import SellerSpriteSettings
    from opscli.seller_sprite.domain.exceptions import SellerSpriteAuthenticationError
    from opscli.seller_sprite.services import api_manager as api_module
    from opscli.seller_sprite.services import task_scheduler as scheduler_module
    from opscli.seller_sprite.services.account_pool import seller_sprite_account_key
    from opscli.seller_sprite.services.task_queue_store import (
        SellerSpriteTaskQueueStore,
    )
    from tests.seller_sprite.test_api_manager import (
        DummyAccountProvider,
        DummyApiClient,
    )

    monkeypatch.setattr(api_module, "SellerSpriteApiClient", DummyApiClient)
    monkeypatch.setattr(api_module, "_upload_export_if_enabled", lambda **kw: None)
    monkeypatch.setattr(
        scheduler_module, "_resolve_task_auth", lambda *a, **kw: (None, None)
    )
    calls = []

    async def browser(**kwargs):
        calls.append(kwargs["payload"]["page"])
        if fail_page_two and kwargs["payload"]["page"] == 2:
            raise SellerSpriteAuthenticationError("模拟第二页登录失效")
        return SimpleNamespace(
            response=response(121, kwargs["payload"]["page"]), warnings=[]
        )

    monkeypatch.setattr(api_module, "_run_browser_route_request", browser)
    store = SellerSpriteTaskQueueStore(db_path=tmp_path / "queue.db")
    provider = DummyAccountProvider()
    account = provider.get_default()
    account_key = seller_sprite_account_key(account)
    scheduler = scheduler_module.SellerSpriteTaskScheduler(
        settings=SellerSpriteSettings(output_dir=tmp_path),
        store=store,
        account_provider=provider,
        auto_start=False,
    )
    request = SellerSpriteScenarioRequest(
        scenario="product-research",
        site="US",
        period="30d",
        job_id="job",
        collection_scope="all",
        mode="browser-route",
    )
    store.enqueue(request=request, queue_scope="test", root_dir=tmp_path / "job")
    claim = store.claim_next_generic_for_account(
        queue_scope="test",
        account_key=account_key,
        assigned_account=account.name,
        worker_key="worker",
        execution_owner=scheduler._session_owner_id,
    )

    async def run():
        if dedicated:
            await scheduler._run_one(
                "job",
                account=account,
                execution_account_key=account_key,
                assignment_generation=claim["assignment_generation"],
            )
        else:
            await scheduler._run_generic_job(
                claimed=claim,
                account=account,
                worker_key="worker",
                attempt_id="attempt",
            )

    asyncio.run(run())
    assert calls == [1, 2]
    status = store.get_status("job")
    assert status["state"] == ("failed" if fail_page_two else "succeeded")
    assert status["failover_count"] == 0
    assert status["collection"]["pages_completed"] == (1 if fail_page_two else 2)
    assert status["collection"]["collection_complete"] is (not fail_page_two)
    assert len(store.load_product_checkpoint("job")["pages"]) == (
        1 if fail_page_two else 2
    )


def test_cli_and_gateway_forward_explicit_scope(monkeypatch):
    from typer.testing import CliRunner

    from opscli.cli import app
    from opscli.mcp.tools import seller_sprite_proxy
    from opscli.seller_sprite import cli

    calls = []
    monkeypatch.setattr(
        cli,
        "SellerSpriteRemoteAdapter",
        lambda: SimpleNamespace(run=lambda **kw: calls.append(kw) or {"success": True}),
    )
    result = CliRunner().invoke(
        app,
        [
            "seller-sprite",
            "run",
            "product-research",
            "--collection-scope",
            "all",
            "--export-format",
            "xlsx",
        ],
    )
    assert result.exit_code == 0
    assert calls[0]["collection_scope"] == "all"

    async def forward(tool, params):
        return params

    monkeypatch.setattr(seller_sprite_proxy, "_call_collector", forward)
    forwarded = asyncio.run(
        seller_sprite_proxy.seller_sprite_run(
            scenario="product-research", collection_scope="all"
        )
    )
    assert forwarded["collection_scope"] == "all"


@pytest.mark.parametrize("entry", ["seller_sprite_run", "seller_sprite_start"])
def test_invalid_all_rejected_before_auth_or_queue(monkeypatch, entry):
    from opscli.mcp.tools import seller_sprite

    async def auth(**kwargs):
        raise AssertionError("无效请求不能触发认证")

    monkeypatch.setattr(seller_sprite, "ensure_ops_credentials", auth)
    monkeypatch.setattr(seller_sprite, "_err", lambda exc, **kw: {"error": str(exc)})
    result = asyncio.run(
        getattr(seller_sprite, entry)(
            scenario="product-research", collection_scope="all", export_format="json"
        )
    )
    assert "xls/xlsx" in result["error"]


def test_all_mcp_bypasses_shared_cache_and_enqueues_scope(monkeypatch):
    from opscli.mcp.tools import seller_sprite
    from opscli.seller_sprite import mcp_bundle

    captured = []

    async def no_cache(*args, **kwargs):
        raise AssertionError("全量模式不能查共享缓存")

    async def auth(**kwargs):
        return None

    async def enqueue(scheduler, request, **kwargs):
        captured.append(request)
        return {"state": "queued", "job_id": "test-all"}

    monkeypatch.setattr(seller_sprite, "find_cached_result", no_cache)
    monkeypatch.setattr(seller_sprite, "ensure_ops_credentials", auth)
    monkeypatch.setattr(seller_sprite, "_enqueue_task_with_auth", enqueue)
    monkeypatch.setattr(seller_sprite, "_get_task_scheduler", lambda: object())
    monkeypatch.setattr(
        seller_sprite, "_get_current_mcp_user_email", lambda: "test@example.com"
    )
    monkeypatch.setattr(
        mcp_bundle, "get_result_cache_context", lambda: (object(), "debug")
    )
    result = asyncio.run(
        seller_sprite.seller_sprite_run(
            scenario="product-research", collection_scope="all"
        )
    )
    assert result["success"] is True
    assert captured[0].collection_scope == "all"


def test_export_recovery_uses_frozen_pages_without_resolving_or_fetching(
    tmp_path, monkeypatch
):
    from opscli.seller_sprite.config import SellerSpriteSettings
    from opscli.seller_sprite.export import xlsx
    from opscli.seller_sprite.services import api_manager as module
    from tests.seller_sprite.test_api_manager import (
        DummyAccountProvider,
        DummyApiClient,
    )

    monkeypatch.setattr(module, "SellerSpriteApiClient", DummyApiClient)
    monkeypatch.setattr(module, "_upload_export_if_enabled", lambda **kw: None)
    saved = []
    calls = []

    async def browser(**kwargs):
        calls.append(kwargs["payload"]["page"])
        return SimpleNamespace(
            response=response(121, kwargs["payload"]["page"]), warnings=[]
        )

    monkeypatch.setattr(module, "_run_browser_route_request", browser)
    manager = module.SellerSpriteApiManager(
        settings=SellerSpriteSettings(output_dir=tmp_path),
        account_provider=DummyAccountProvider(),
    )
    manager.product_task_root = tmp_path
    manager.product_checkpoint_listener = lambda value: saved.append(
        json.loads(json.dumps(value))
    )
    request = SellerSpriteScenarioRequest(
        scenario="product-research",
        site="US",
        period="30d",
        job_id="job",
        mode="browser-route",
        collection_scope="all",
        attempt_output_dir=str(tmp_path / "first"),
    )
    export = xlsx.export_product_pages_to_xlsx

    def fail_export(**kwargs):
        raise OSError("模拟磁盘写入失败")

    monkeypatch.setattr(xlsx, "export_product_pages_to_xlsx", fail_export)
    with pytest.raises(OSError, match="磁盘"):
        asyncio.run(manager.run(request))
    assert calls == [1, 2]
    manager.product_checkpoint = saved[-1]

    async def no_resolve(*args, **kwargs):
        raise AssertionError("恢复不能重算月份或类目")

    monkeypatch.setattr(
        module.SellerSpriteCategoryResolver, "resolve_params", no_resolve
    )
    monkeypatch.setattr(xlsx, "export_product_pages_to_xlsx", export)
    result = asyncio.run(
        manager.run(replace(request, attempt_output_dir=str(tmp_path / "second")))
    )
    assert result.row_count == 121
    assert calls == [1, 2]
    assert all("first/pages" in page["path"] for page in saved[-1]["pages"])


def test_page_size_guard_stops_before_page_commit(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "opscli.seller_sprite.services.product_pagination.MAX_PAGE_BYTES", 32
    )
    instance = collector(tmp_path)

    async def fetch(payload, root):
        return response(1, 1)

    with pytest.raises(SellerSpriteConfigError, match="响应超过"):
        asyncio.run(instance.collect(fetch))
    assert instance.state["pages"] == []
    assert instance.summary()["current_page"] == 1
    assert instance.summary()["collection_complete"] is False


def test_empty_streaming_export_contains_headers(tmp_path):
    from opscli.seller_sprite.export.xlsx import export_product_pages_to_xlsx

    result = export_product_pages_to_xlsx(
        rows=iter(()),
        output_path=tmp_path / "empty.xlsx",
        site="US",
        period="30d",
        params={},
    )
    workbook = load_workbook(result.path, read_only=True)
    try:
        rows = list(workbook.active.values)
        assert len(rows) == 1
        assert rows[0][0] == "ASIN"
    finally:
        workbook.close()
